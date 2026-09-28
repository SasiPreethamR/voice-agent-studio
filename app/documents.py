from __future__ import annotations

import csv
import hashlib
import importlib.machinery
import io
import json
import math
import mimetypes
import re
import shutil
import sqlite3
import sys
import time
import types
import uuid
import zipfile
from pathlib import Path, PurePosixPath
from typing import Optional
from xml.etree import ElementTree as ET

import numpy as np

from app.config import DATA_DIR
from app.providers import DEFAULT_MODEL, ModelSettings, chat_completion

DOCS_DIR = DATA_DIR / "documents"
RAW_DIR = DOCS_DIR / "raw"
TREE_DIR = DOCS_DIR / "tree"
ARTIFACTS_DIR = DOCS_DIR / "artifacts"
DB_DIR = DOCS_DIR / "db"
DB_PATH = DB_DIR / "documents.sqlite"
EMBED_DIM = 384
MIN_RETRIEVAL_CONFIDENCE = 0.60
MAX_VIEWER_CHARS = 500000
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
ARCHIVE_EXTS = {".zip"}


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def _new_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def _safe_name(name: str) -> str:
    name = (name or "document").replace("\\", "/").split("/")[-1]
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .") or "document"


def _safe_relative_path(path: str) -> str:
    cleaned = (path or "document").replace("\\", "/").lstrip("/")
    parts = []
    for part in PurePosixPath(cleaned).parts:
        if part in {"", ".", ".."}:
            continue
        parts.append(re.sub(r"[^A-Za-z0-9._ ()\[\]{}@,+='&-]+", "_", part).strip(" .") or "item")
    return str(PurePosixPath(*parts)) if parts else "document"


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except Exception:
            pass
    return content.decode("utf-8", errors="ignore")


def _token_count(text: str) -> int:
    return len(re.findall(r"[\w\u0900-\u0d7f']+", text or ""))


class HashingEmbedder:
    dim = EMBED_DIM

    def encode(self, texts, show_progress_bar: bool = False):
        vectors = []
        for text in texts:
            vec = np.zeros(self.dim, dtype=np.float32)
            words = re.findall(r"[\w\u0900-\u0d7f']+", (text or "").lower())
            for word in words:
                digest = hashlib.blake2b(word.encode("utf-8", errors="ignore"), digest_size=8).digest()
                value = int.from_bytes(digest, "little", signed=False)
                sign = 1.0 if (value >> 63) else -1.0
                vec[value % self.dim] += sign * (1.0 + min(len(word), 18) / 18.0)
            compact = re.sub(r"\s+", " ", (text or "").lower())
            for index in range(max(0, len(compact) - 4)):
                digest = hashlib.blake2b(compact[index:index + 5].encode("utf-8", errors="ignore"), digest_size=8).digest()
                value = int.from_bytes(digest, "little", signed=False)
                vec[value % self.dim] += 0.12 if (value >> 63) else -0.12
            norm = float(np.linalg.norm(vec))
            if norm > 0:
                vec /= norm
            vectors.append(vec)
        return np.vstack(vectors) if vectors else np.zeros((0, self.dim), dtype=np.float32)


class EmbeddingBackend:
    def __init__(self):
        self.model = None
        self.fallback = False

    def _load(self):
        if self.model is not None:
            return
        try:
            self._shim_torchcodec_for_text_embeddings()
            from sentence_transformers import SentenceTransformer
            print("[DOCS] Loading embedding model all-MiniLM-L6-v2 ...")
            self.model = SentenceTransformer("all-MiniLM-L6-v2")
            self.fallback = False
            print("[DOCS] Embedding model ready.")
        except Exception as exc:
            reason = str(exc).splitlines()[0]
            print(f"[DOCS] sentence-transformers unavailable ({type(exc).__name__}: {reason}); using hashing embedder fallback.")
            self.model = HashingEmbedder()
            self.fallback = True

    @staticmethod
    def _shim_torchcodec_for_text_embeddings():
        if "torchcodec.decoders" in sys.modules:
            return

        class _UnavailableDecoder:
            pass

        torchcodec = types.ModuleType("torchcodec")
        decoders = types.ModuleType("torchcodec.decoders")
        torchcodec.__spec__ = importlib.machinery.ModuleSpec("torchcodec", loader=None)
        decoders.__spec__ = importlib.machinery.ModuleSpec("torchcodec.decoders", loader=None)
        decoders.AudioDecoder = _UnavailableDecoder
        decoders.VideoDecoder = _UnavailableDecoder
        torchcodec.decoders = decoders
        sys.modules["torchcodec"] = torchcodec
        sys.modules["torchcodec.decoders"] = decoders

    def encode(self, texts):
        self._load()
        if not texts:
            return np.zeros((0, EMBED_DIM), dtype=np.float32)
        vectors = np.asarray(self.model.encode(texts, show_progress_bar=False), dtype=np.float32)
        if vectors.ndim == 1:
            vectors = vectors.reshape(1, -1)
        return vectors / (np.linalg.norm(vectors, axis=1, keepdims=True) + 1e-8)


class DocumentStore:
    def __init__(self):
        for path in (DOCS_DIR, RAW_DIR, TREE_DIR, ARTIFACTS_DIR, DB_DIR):
            path.mkdir(parents=True, exist_ok=True)
        self.conn = self._open_db()
        self.embedder = EmbeddingBackend()
        self._cache = {}
        self.fts_enabled = False
        self._init_db()

    @staticmethod
    def _open_db() -> sqlite3.Connection:
        """Open the index; a corrupt DB is moved aside and rebuilt from artifacts/."""
        conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        try:
            result = conn.execute("PRAGMA quick_check").fetchone()[0]
        except sqlite3.DatabaseError as exc:
            result = str(exc)
        if result == "ok":
            return conn
        conn.close()
        stamp = time.strftime("%Y%m%d-%H%M%S")
        for suffix in ("", "-wal", "-shm"):
            src = DB_PATH.with_name(DB_PATH.name + suffix)
            if src.exists():
                src.rename(src.with_name(f"{DB_PATH.name}.corrupt-{stamp}{suffix}"))
        print(f"[DOCS] Document index was corrupt ({result}); moved it aside as "
              f"{DB_PATH.name}.corrupt-{stamp} and rebuilding from extracted artifacts.")
        conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self):
        self.conn.executescript(
            """
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS collections(
                id TEXT PRIMARY KEY,
                name TEXT,
                status TEXT,
                total_files INTEGER DEFAULT 0,
                ready_files INTEGER DEFAULT 0,
                failed_files INTEGER DEFAULT 0,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS jobs(
                id TEXT PRIMARY KEY,
                collection_id TEXT,
                status TEXT,
                total_files INTEGER DEFAULT 0,
                processed_files INTEGER DEFAULT 0,
                failed_files INTEGER DEFAULT 0,
                quest_total_files INTEGER DEFAULT 0,
                quest_processed_files INTEGER DEFAULT 0,
                quest_count INTEGER DEFAULT 0,
                error TEXT,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS files(
                id TEXT PRIMARY KEY,
                collection_id TEXT,
                sha256 TEXT,
                filename TEXT,
                relative_path TEXT,
                mime_type TEXT,
                extension TEXT,
                size_bytes INTEGER,
                parser TEXT,
                status TEXT,
                error TEXT,
                char_count INTEGER DEFAULT 0,
                token_count INTEGER DEFAULT 0,
                chunk_count INTEGER DEFAULT 0,
                quest_count INTEGER DEFAULT 0,
                quest_status TEXT DEFAULT 'pending',
                raw_path TEXT,
                artifact_path TEXT,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS chunks(
                id TEXT PRIMARY KEY,
                file_id TEXT,
                collection_id TEXT,
                chunk_index INTEGER,
                chunk_type TEXT,
                text TEXT,
                source_path TEXT,
                page INTEGER,
                sheet TEXT,
                row_start INTEGER,
                row_end INTEGER,
                token_count INTEGER DEFAULT 0,
                metadata_json TEXT
            );
            CREATE TABLE IF NOT EXISTS quests(
                id TEXT PRIMARY KEY,
                file_id TEXT,
                collection_id TEXT,
                chunk_id TEXT,
                question TEXT,
                answer TEXT,
                score REAL DEFAULT 1.0,
                metadata_json TEXT
            );
            CREATE TABLE IF NOT EXISTS embeddings(
                owner_type TEXT,
                owner_id TEXT,
                dim INTEGER,
                vector BLOB,
                PRIMARY KEY(owner_type, owner_id)
            );
            CREATE INDEX IF NOT EXISTS idx_files_collection ON files(collection_id);
            CREATE INDEX IF NOT EXISTS idx_files_path ON files(relative_path);
            CREATE INDEX IF NOT EXISTS idx_chunks_file ON chunks(file_id);
            CREATE INDEX IF NOT EXISTS idx_quests_file ON quests(file_id);
            """
        )
        self.conn.commit()
        self._ensure_schema_columns()
        self._init_fts()
        self._rehydrate_from_artifacts_if_empty()
        self._ensure_fts_index()

    def _ensure_schema_columns(self):
        columns = {row["name"] for row in self.conn.execute("PRAGMA table_info(jobs)").fetchall()}
        for name in ("quest_total_files", "quest_processed_files", "quest_count"):
            if name not in columns:
                self.conn.execute(f"ALTER TABLE jobs ADD COLUMN {name} INTEGER DEFAULT 0")
        self.conn.commit()

    @property
    def documents(self) -> dict:
        return {doc["id"]: doc for doc in self.list_documents()}

    def _init_fts(self):
        for tokenizer in ("porter unicode61", "unicode61"):
            try:
                self.conn.execute(
                    f"""CREATE VIRTUAL TABLE IF NOT EXISTS chunk_fts
                    USING fts5(chunk_id UNINDEXED, file_id UNINDEXED, source_path, text, tokenize='{tokenizer}')"""
                )
                self.conn.commit()
                self.fts_enabled = True
                return
            except sqlite3.OperationalError:
                continue
        print("[DOCS] SQLite FTS5 unavailable; chunk full-text search disabled.")
        self.fts_enabled = False

    def _ensure_fts_index(self):
        if not self.fts_enabled:
            return
        try:
            indexed = self.conn.execute("SELECT COUNT(*) FROM chunk_fts").fetchone()[0]
            total = self.conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
            if indexed == total:
                return
            self.conn.execute("DELETE FROM chunk_fts")
            rows = self.conn.execute("SELECT id, file_id, source_path, text FROM chunks").fetchall()
            for row in rows:
                self._put_chunk_fts(row["id"], row["file_id"], row["source_path"], row["text"])
            self.conn.commit()
            print(f"[DOCS] Rebuilt FTS index for {len(rows)} chunks.")
        except sqlite3.Error as exc:
            print(f"[DOCS] FTS index rebuild failed: {exc}")

    def _put_chunk_fts(self, chunk_id: str, file_id: str, source_path: str, text: str):
        if not self.fts_enabled:
            return
        self.conn.execute("DELETE FROM chunk_fts WHERE chunk_id=?", (chunk_id,))
        self.conn.execute(
            "INSERT INTO chunk_fts(chunk_id, file_id, source_path, text) VALUES(?, ?, ?, ?)",
            (chunk_id, file_id, source_path or "", text or ""),
        )

    def _rehydrate_from_artifacts_if_empty(self):
        existing = self.conn.execute("SELECT COUNT(*) FROM files").fetchone()[0]
        artifacts = sorted(ARTIFACTS_DIR.glob("file_*/extracted.json"))
        if existing or not artifacts:
            return
        print(f"[DOCS] Rehydrating {len(artifacts)} extracted documents into SQLite index...")
        collections: dict[str, str] = {}
        restored = failed = 0
        for artifact_path in artifacts:
            try:
                data = json.loads(artifact_path.read_text(encoding="utf-8"))
                rel_path = _safe_relative_path(data.get("relative_path") or data.get("filename") or artifact_path.parent.name)
                filename = _safe_name(data.get("filename") or PurePosixPath(rel_path).name)
                file_id = data.get("file_id") or artifact_path.parent.name or _new_id("file")
                root_name = rel_path.split("/", 1)[0] if "/" in rel_path else "Documents"
                collection_id = collections.get(root_name)
                if not collection_id:
                    collection_id = self.create_collection(root_name, 0)
                    collections[root_name] = collection_id
                elements = data.get("elements") or []
                full_text = "\n\n".join(element.get("text", "") for element in elements if element.get("text", "")).strip()
                if not full_text:
                    text_file = artifact_path.parent / "text.md"
                    full_text = text_file.read_text(encoding="utf-8") if text_file.exists() else ""
                    elements = [{"type": "text", "text": full_text, "source_path": rel_path}] if full_text else []
                if not full_text:
                    failed += 1
                    continue
                raw_source = self._find_raw_source(filename)
                raw_bytes = raw_source.read_bytes() if raw_source and raw_source.exists() else full_text.encode("utf-8")
                sha = hashlib.sha256(raw_bytes).hexdigest()
                tree_path = TREE_DIR / rel_path
                tree_path.parent.mkdir(parents=True, exist_ok=True)
                if not tree_path.exists():
                    tree_path.write_bytes(raw_bytes)
                chunks = self._chunk_elements(elements)
                if not chunks:
                    failed += 1
                    continue
                now = _now()
                self.conn.execute(
                    """INSERT OR REPLACE INTO files(id,collection_id,sha256,filename,relative_path,mime_type,extension,size_bytes,parser,status,error,char_count,token_count,chunk_count,quest_count,quest_status,raw_path,artifact_path,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        file_id, collection_id, sha, filename, rel_path,
                        mimetypes.guess_type(filename)[0] or "application/octet-stream", Path(filename).suffix.lower(),
                        len(raw_bytes), data.get("parser") or "artifact", "ready", None, len(full_text),
                        _token_count(full_text), len(chunks), 0, "pending", str(tree_path), str(artifact_path), now, now,
                    ),
                )
                self.conn.execute("DELETE FROM chunks WHERE file_id=?", (file_id,))
                self.conn.execute("DELETE FROM quests WHERE file_id=?", (file_id,))
                self.conn.execute("DELETE FROM embeddings WHERE owner_type='doc' AND owner_id=?", (file_id,))
                chunk_vectors = self.embedder.encode([chunk["text"] for chunk in chunks])
                doc_vector = self.embedder.encode([self._document_embedding_text(filename, rel_path, full_text, chunks)])[0]
                for index, chunk in enumerate(chunks):
                    chunk_id = _new_id("chunk")
                    chunk["id"] = chunk_id
                    self.conn.execute(
                        """INSERT INTO chunks(id,file_id,collection_id,chunk_index,chunk_type,text,source_path,page,sheet,row_start,row_end,token_count,metadata_json)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            chunk_id, file_id, collection_id, index, chunk.get("type", "chunk"), chunk["text"],
                            chunk.get("source_path", rel_path), chunk.get("page"), chunk.get("sheet"),
                            chunk.get("row_start"), chunk.get("row_end"), _token_count(chunk["text"]),
                            json.dumps(chunk.get("metadata", {}), ensure_ascii=False),
                        ),
                    )
                    self._put_embedding("chunk", chunk_id, chunk_vectors[index])
                    self._put_chunk_fts(chunk_id, file_id, chunk.get("source_path", rel_path), chunk["text"])
                self._put_embedding("doc", file_id, doc_vector)
                self._generate_heuristic_quests(file_id)
                self.conn.commit()
                restored += 1
            except Exception as exc:
                failed += 1
                print(f"[DOCS] Artifact rehydrate failed for {artifact_path}: {exc}")
        for collection_id in collections.values():
            self._update_collection_counts(collection_id)
        self._invalidate_cache()
        print(f"[DOCS] Rehydrated {restored} documents ({failed} failed).")

    def _find_raw_source(self, filename: str) -> Optional[Path]:
        for path in RAW_DIR.glob(f"*/{filename}"):
            if path.is_file():
                return path
        return None

    def _generate_heuristic_quests(self, file_id: str) -> int:
        file_row = self.conn.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
        if not file_row:
            return 0
        chunks = self.conn.execute("SELECT * FROM chunks WHERE file_id=? ORDER BY chunk_index", (file_id,)).fetchall()
        if not chunks:
            return 0
        target = min(8, self._quest_target_count(max(file_row["token_count"] or 0, sum(chunk["token_count"] or 0 for chunk in chunks))))
        selected = self._select_quest_chunks(chunks, target)
        quests = []
        seen = set()
        for chunk in selected:
            needed = max(1, min(3, target - len(quests)))
            for question in self._fallback_questions(file_row, chunk, needed):
                cleaned = self._clean_question(question)
                key = cleaned.lower()
                if cleaned and key not in seen:
                    seen.add(key)
                    quests.append({"question": cleaned, "chunk_id": chunk["id"], "answer": chunk["text"][:360]})
                if len(quests) >= target:
                    break
            if len(quests) >= target:
                break
        if not quests:
            return 0
        vectors = self.embedder.encode([quest["question"] for quest in quests])
        self.conn.execute("DELETE FROM quests WHERE file_id=?", (file_id,))
        for quest, vector in zip(quests, vectors):
            quest_id = _new_id("quest")
            self.conn.execute(
                "INSERT INTO quests(id,file_id,collection_id,chunk_id,question,answer,metadata_json) VALUES(?,?,?,?,?,?,?)",
                (quest_id, file_id, file_row["collection_id"], quest["chunk_id"], quest["question"], quest["answer"], "{}"),
            )
            self._put_embedding("quest", quest_id, vector)
        self.conn.execute("UPDATE files SET quest_status='done',quest_count=?,updated_at=? WHERE id=?", (len(quests), _now(), file_id))
        return len(quests)

    def create_collection(self, name: Optional[str] = None, total_files: int = 0) -> str:
        collection_id = _new_id("col")
        now = _now()
        self.conn.execute(
            "INSERT INTO collections(id,name,status,total_files,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (collection_id, name or "Upload", "ingesting", total_files, now, now),
        )
        self.conn.commit()
        return collection_id

    def create_job(self, collection_id: str, total_files: int) -> str:
        job_id = _new_id("job")
        now = _now()
        self.conn.execute(
            "INSERT INTO jobs(id,collection_id,status,total_files,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (job_id, collection_id, "queued", total_files, now, now),
        )
        self.conn.commit()
        return job_id

    def get_job(self, job_id: str) -> Optional[dict]:
        row = self.conn.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        return dict(row) if row else None

    def list_jobs(self) -> list:
        rows = self.conn.execute(
            "SELECT * FROM jobs ORDER BY created_at DESC LIMIT 200"
        ).fetchall()
        return [dict(r) for r in rows]

    def retry_collection(self, collection_id: str) -> Optional[str]:
        """Create a retry job that covers:
        - files with status='failed'  (will be re-parsed from the tree directory on disk)
        - files stuck in quest_status='generating' (server crash mid-K-Quest; reset to pending)
        - files with quest_status='pending'  (K-Quest never ran)
        Returns job_id, or None if there is nothing to retry.
        """
        # Files where parsing itself failed — raw bytes still on disk under TREE_DIR
        failed_files = self.conn.execute(
            "SELECT id FROM files WHERE collection_id=? AND status='failed'",
            (collection_id,),
        ).fetchall()

        # Files whose K-Quest got stuck mid-run (process died): reset them to pending
        stuck = self.conn.execute(
            "SELECT id FROM files WHERE collection_id=? AND status='ready' AND quest_status='generating'",
            (collection_id,),
        ).fetchall()
        if stuck:
            self.conn.execute(
                "UPDATE files SET quest_status='pending', updated_at=? "
                "WHERE collection_id=? AND status='ready' AND quest_status='generating'",
                (_now(), collection_id),
            )
            self.conn.commit()

        # Files with quest_status still pending (includes the ones we just reset above)
        pending_quest = self.conn.execute(
            "SELECT id FROM files WHERE collection_id=? AND status='ready' "
            "AND (quest_status IS NULL OR quest_status='pending')",
            (collection_id,),
        ).fetchall()

        total = len(failed_files) + len(pending_quest)
        if total == 0:
            return None
        job_id = self.create_job(collection_id, total)
        return job_id

    async def run_retry_quests(self, job_id: str, collection_id: str, model: ModelSettings = DEFAULT_MODEL):
        """Full retry: re-parse failed files from disk, then run K-Quest for all pending files."""
        import asyncio

        job = self.get_job(job_id)
        if not job:
            return

        # ── Step 1: Re-parse files whose parsing previously failed ──────────
        failed_files = self.conn.execute(
            "SELECT id, filename, relative_path FROM files WHERE collection_id=? AND status='failed'",
            (collection_id,),
        ).fetchall()

        parse_processed = parse_failed = 0
        if failed_files:
            self.update_job(job_id, status="running", processed_files=0, failed_files=0)
            for row in failed_files:
                try:
                    # The raw file bytes are always written to TREE_DIR before parsing runs
                    tree_path = TREE_DIR / row["relative_path"]
                    if tree_path.exists():
                        content = tree_path.read_bytes()
                        self.delete_document(row["id"])  # remove the failed DB entry
                        await asyncio.to_thread(
                            self.add_document,
                            row["filename"],
                            content,
                            row["relative_path"],
                            collection_id,
                        )
                    else:
                        print(f"[DOCS] Retry: source file not found on disk: {row['relative_path']}")
                        parse_failed += 1
                except Exception as exc:
                    print(f"[DOCS] Retry re-parse failed for {row['relative_path']}: {exc}")
                    parse_failed += 1
                parse_processed += 1
                self.update_job(job_id, processed_files=parse_processed, failed_files=parse_failed)

        # ── Step 2: K-Quest for all ready files that still need it ──────────
        pending_quest = self.conn.execute(
            "SELECT id FROM files WHERE collection_id=? AND status='ready' "
            "AND (quest_status IS NULL OR quest_status='pending')",
            (collection_id,),
        ).fetchall()

        self.update_job(
            job_id,
            status="generating_quests",
            processed_files=parse_processed,
            quest_total_files=len(pending_quest),
            quest_processed_files=0,
            quest_count=0,
        )
        quest_processed = quest_count = 0
        for row in pending_quest:
            try:
                quest_count += await self.generate_quests(row["id"], model=model) or 0
            except Exception as exc:
                print(f"[DOCS] Retry quest failed for {row['id']}: {exc}")
            quest_processed += 1
            self.update_job(job_id, quest_processed_files=quest_processed, quest_count=quest_count)

        self.update_job(job_id, status="done" if parse_failed == 0 else "done_with_errors")
        self._update_collection_counts(collection_id)

    def update_job(self, job_id: str, **fields):
        if not fields:
            return
        fields["updated_at"] = _now()
        sets = ", ".join(f"{key}=?" for key in fields)
        self.conn.execute(f"UPDATE jobs SET {sets} WHERE id=?", (*fields.values(), job_id))
        self.conn.commit()

    def add_document(self, filename: str, content_bytes: bytes, relative_path: Optional[str] = None,
                     collection_id: Optional[str] = None) -> str:
        if not content_bytes:
            raise ValueError("Empty file")
        collection_id = collection_id or self.create_collection(_safe_name(filename), 1)
        rel_path = _safe_relative_path(relative_path or filename or "document")
        filename = _safe_name(filename or PurePosixPath(rel_path).name)
        file_id = _new_id("file")
        sha = hashlib.sha256(content_bytes).hexdigest()
        ext = Path(filename).suffix.lower()
        mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"

        raw_dir = RAW_DIR / sha
        raw_dir.mkdir(parents=True, exist_ok=True)
        raw_path = raw_dir / filename
        if not raw_path.exists():
            raw_path.write_bytes(content_bytes)
        tree_path = TREE_DIR / rel_path
        tree_path.parent.mkdir(parents=True, exist_ok=True)
        tree_path.write_bytes(content_bytes)

        now = _now()
        self.conn.execute(
            """INSERT INTO files(id,collection_id,sha256,filename,relative_path,mime_type,extension,size_bytes,status,raw_path,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (file_id, collection_id, sha, filename, rel_path, mime_type, ext, len(content_bytes), "parsing", str(tree_path), now, now),
        )
        self.conn.commit()
        try:
            parsed = self._parse_file(filename, rel_path, content_bytes)
            elements = parsed["elements"]
            full_text = "\n\n".join(element.get("text", "") for element in elements if element.get("text", "")).strip()
            if not full_text:
                raise ValueError(parsed.get("error") or "Could not extract text from document")
            chunks = self._chunk_elements(elements)
            if not chunks:
                raise ValueError("Could not create chunks from document")
            artifact_dir = ARTIFACTS_DIR / file_id
            artifact_dir.mkdir(parents=True, exist_ok=True)
            artifact_path = artifact_dir / "extracted.json"
            artifact_path.write_text(json.dumps({
                "file_id": file_id,
                "filename": filename,
                "relative_path": rel_path,
                "parser": parsed["parser"],
                "elements": elements,
                "text_preview": full_text[:4000],
            }, ensure_ascii=False, indent=2), encoding="utf-8")
            (artifact_dir / "text.md").write_text(full_text, encoding="utf-8")

            chunk_vectors = self.embedder.encode([chunk["text"] for chunk in chunks])
            doc_vector = self.embedder.encode([self._document_embedding_text(filename, rel_path, full_text, chunks)])[0]
            for index, chunk in enumerate(chunks):
                chunk_id = _new_id("chunk")
                chunk["id"] = chunk_id
                self.conn.execute(
                    """INSERT INTO chunks(id,file_id,collection_id,chunk_index,chunk_type,text,source_path,page,sheet,row_start,row_end,token_count,metadata_json)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (chunk_id, file_id, collection_id, index, chunk.get("type", "chunk"), chunk["text"],
                     chunk.get("source_path", rel_path), chunk.get("page"), chunk.get("sheet"),
                     chunk.get("row_start"), chunk.get("row_end"), _token_count(chunk["text"]),
                     json.dumps(chunk.get("metadata", {}), ensure_ascii=False)),
                )
                self._put_embedding("chunk", chunk_id, chunk_vectors[index])
                self._put_chunk_fts(chunk_id, file_id, chunk.get("source_path", rel_path), chunk["text"])
            self._put_embedding("doc", file_id, doc_vector)
            self.conn.execute(
                """UPDATE files SET parser=?,status='ready',error=NULL,char_count=?,token_count=?,chunk_count=?,artifact_path=?,updated_at=?
                   WHERE id=?""",
                (parsed["parser"], len(full_text), _token_count(full_text), len(chunks), str(artifact_path), _now(), file_id),
            )
            self.conn.commit()
            self._invalidate_cache()
            self._update_collection_counts(collection_id)
            return file_id
        except Exception as exc:
            self.conn.execute("UPDATE files SET status='failed',error=?,updated_at=? WHERE id=?", (str(exc), _now(), file_id))
            self.conn.commit()
            self._update_collection_counts(collection_id)
            raise

    async def process_ingest_job(self, job_id: str, files: list[dict], model: ModelSettings = DEFAULT_MODEL):
        import asyncio

        job = self.get_job(job_id)
        if not job:
            return
        collection_id = job["collection_id"]
        self.update_job(job_id, status="running")
        processed = failed = 0
        for item in files:
            try:
                await asyncio.to_thread(self.add_document, item["filename"], item["content"], item.get("relative_path"), collection_id)
            except Exception as exc:
                print(f"[DOCS] Ingest failed for {item.get('relative_path') or item.get('filename')}: {exc}")
                failed += 1
            processed += 1
            self.update_job(job_id, processed_files=processed, failed_files=failed)
        quest_rows = self.conn.execute("SELECT id FROM files WHERE collection_id=? AND status='ready'", (collection_id,)).fetchall()
        self.update_job(job_id, status="generating_quests", quest_total_files=len(quest_rows), quest_processed_files=0, quest_count=0)
        quest_processed = quest_count = 0
        for row in quest_rows:
            try:
                quest_count += await self.generate_quests(row["id"], model=model) or 0
            except Exception as exc:
                print(f"[DOCS] Quest generation failed for {row['id']}: {exc}")
            quest_processed += 1
            self.update_job(job_id, quest_processed_files=quest_processed, quest_count=quest_count)
        self.update_job(job_id, status="done" if failed == 0 else "done_with_errors")
        self._update_collection_counts(collection_id)

    async def generate_quests(self, file_id: str, model: ModelSettings = DEFAULT_MODEL):
        file_row = self.conn.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
        if not file_row or file_row["status"] != "ready":
            return 0
        self.conn.execute("UPDATE files SET quest_status='generating',updated_at=? WHERE id=?", (_now(), file_id))
        self.conn.commit()
        chunks = self.conn.execute("SELECT * FROM chunks WHERE file_id=? ORDER BY chunk_index", (file_id,)).fetchall()
        if not chunks:
            return 0
        target_k = self._quest_target_count(max(file_row["token_count"] or 0, sum(chunk["token_count"] or 0 for chunk in chunks)))
        selected = self._select_quest_chunks(chunks, target_k)
        quests = []
        for chunk in selected:
            remaining = target_k - len(quests)
            if remaining <= 0:
                break
            needed = min(4, remaining)
            prompt = (
                "Generate likely user questions that the following document excerpt can answer.\n"
                "Return one question per line. Do not include numbering, answers, markdown, or commentary.\n"
                f"Document: {file_row['relative_path']}\nQuestions to generate: {needed}\n\nExcerpt:\n{chunk['text'][:3500]}"
            )
            try:
                content = await chat_completion(
                    [{"role": "user", "content": prompt}], model=model, max_tokens=320, temperature=0.25,
                )
                questions = self._parse_questions(content)
            except Exception as exc:
                print(f"[DOCS] LLM quest generation failed for {file_id}: {exc}")
                questions = self._fallback_questions(file_row, chunk, needed)
            seen = {quest["question"].lower() for quest in quests}
            for question in questions:
                cleaned = self._clean_question(question)
                if cleaned and cleaned.lower() not in seen:
                    quests.append({"question": cleaned, "chunk_id": chunk["id"], "answer": chunk["text"][:360]})
                    seen.add(cleaned.lower())
                if len(quests) >= target_k:
                    break
        old_ids = [row["id"] for row in self.conn.execute("SELECT id FROM quests WHERE file_id=?", (file_id,)).fetchall()]
        self.conn.execute("DELETE FROM quests WHERE file_id=?", (file_id,))
        for quest_id in old_ids:
            self.conn.execute("DELETE FROM embeddings WHERE owner_type='quest' AND owner_id=?", (quest_id,))
        if quests:
            vectors = self.embedder.encode([quest["question"] for quest in quests])
            for quest, vector in zip(quests, vectors):
                quest_id = _new_id("quest")
                self.conn.execute(
                    "INSERT INTO quests(id,file_id,collection_id,chunk_id,question,answer,metadata_json) VALUES(?,?,?,?,?,?,?)",
                    (quest_id, file_id, file_row["collection_id"], quest.get("chunk_id"), quest["question"], quest.get("answer", ""), "{}"),
                )
                self._put_embedding("quest", quest_id, vector)
        self.conn.execute("UPDATE files SET quest_status='done',quest_count=?,updated_at=? WHERE id=?", (len(quests), _now(), file_id))
        self.conn.commit()
        self._invalidate_cache()
        print(f"[DOCS] Generated {len(quests)} quests for {file_row['relative_path']}")
        return len(quests)

    generate_faqs = generate_quests

    def retrieve(self, query: str, top_k: int = 5, scopes: Optional[list[dict]] = None) -> list:
        if not query.strip():
            return []
        q_vec = self.embedder.encode([query])[0]
        scope = self._normalize_scopes(scopes)
        candidates = {}
        for hit in self._fts_search(query, max(20, top_k * 5), scope):
            chunk = self._get_chunk(hit["id"])
            if chunk:
                self._merge_candidate(candidates, chunk["id"], hit["score"], chunk["text"], chunk["file_id"], "fts", "", chunk)
        for hit in self._search("quest", q_vec, max(20, top_k * 5), scope):
            quest = self._get_quest(hit["id"])
            if not quest:
                continue
            chunk = self._get_chunk(quest["chunk_id"]) if quest["chunk_id"] else None
            key = chunk["id"] if chunk else f"quest:{quest['id']}"
            prefix = f"Matched question: {quest['question']}\n"
            self._merge_candidate(candidates, key, hit["score"], chunk["text"] if chunk else quest["answer"], quest["file_id"], "quest", prefix, chunk)
        for hit in self._search("doc", q_vec, max(5, top_k * 2), scope):
            file_row = self._get_file(hit["id"])
            if not file_row:
                continue
            chunk = self._first_chunk(file_row["id"])
            self._merge_candidate(candidates, f"doc:{file_row['id']}", hit["score"] * 0.72, chunk["text"] if chunk else file_row["relative_path"], file_row["id"], "doc", "", chunk)
        for hit in self._search("chunk", q_vec, max(10, top_k * 3), scope):
            chunk = self._get_chunk(hit["id"])
            if chunk:
                self._merge_candidate(candidates, chunk["id"], hit["score"] * 0.64, chunk["text"], chunk["file_id"], "chunk", "", chunk)
        ranked = sorted(candidates.values(), key=lambda item: item["score"], reverse=True)
        if not ranked or ranked[0]["score"] < MIN_RETRIEVAL_CONFIDENCE:
            return []
        floor = max(MIN_RETRIEVAL_CONFIDENCE, ranked[0]["score"] * 0.72)
        return [item for item in ranked if item["score"] >= floor][:top_k]

    def list_documents(self) -> list:
        rows = self.conn.execute(
            """SELECT f.*, c.name AS collection_name FROM files f
               LEFT JOIN collections c ON c.id=f.collection_id
               ORDER BY f.created_at DESC"""
        ).fetchall()
        return [self._file_to_dict(row) for row in rows]

    def list_collections(self) -> list:
        rows = self.conn.execute(
            "SELECT * FROM collections WHERE total_files > 0 OR ready_files > 0 OR failed_files > 0 ORDER BY created_at DESC"
        ).fetchall()
        return [dict(row) for row in rows]

    def list_tree(self) -> dict:
        root = {"name": "Documents", "path": "", "type": "root", "children": {}, "files": []}
        for doc in self.list_documents():
            path = doc.get("relative_path") or doc.get("filename") or "document"
            parts = [part for part in PurePosixPath(path).parts if part]
            node = root
            for index, part in enumerate(parts):
                if index == len(parts) - 1:
                    node["files"].append(doc)
                else:
                    folder_path = "/".join(parts[:index + 1])
                    node = node["children"].setdefault(part, {"name": part, "path": folder_path, "type": "folder", "children": {}, "files": []})
        return self._tree_to_public(root)

    def get_document_view(self, file_id: str) -> Optional[dict]:
        row = self.conn.execute(
            """SELECT f.*, c.name AS collection_name FROM files f
               LEFT JOIN collections c ON c.id=f.collection_id
               WHERE f.id=?""",
            (file_id,),
        ).fetchone()
        if not row:
            return None
        chunks = [dict(chunk) for chunk in self.conn.execute(
            "SELECT id,chunk_index,chunk_type,text,source_path,page,sheet,row_start,row_end,token_count FROM chunks WHERE file_id=? ORDER BY chunk_index",
            (file_id,),
        ).fetchall()]
        quests = [dict(quest) for quest in self.conn.execute(
            "SELECT id,chunk_id,question,answer,score FROM quests WHERE file_id=? ORDER BY rowid",
            (file_id,),
        ).fetchall()]
        text = self._document_view_text(row, chunks)
        truncated = len(text) > MAX_VIEWER_CHARS
        if truncated:
            text = text[:MAX_VIEWER_CHARS]
        return {
            "document": self._file_to_dict(row),
            "text": text,
            "truncated": truncated,
            "chunks": chunks,
            "quests": quests,
        }

    def _document_view_text(self, file_row, chunks: list[dict]) -> str:
        artifact_path = Path(file_row["artifact_path"] or "")
        text_path = artifact_path.parent / "text.md" if artifact_path.name else Path("")
        try:
            if text_path.exists() and ARTIFACTS_DIR.resolve() in text_path.resolve().parents:
                return text_path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            pass
        return "\n\n".join(chunk.get("text") or "" for chunk in chunks).strip()

    def delete_document(self, file_id: str) -> bool:
        row = self._get_file(file_id)
        if not row:
            return False
        collection_id = row["collection_id"]
        mirror_path = Path(row["raw_path"] or "")
        artifact_path = Path(row["artifact_path"] or "")
        chunk_ids = [item["id"] for item in self.conn.execute("SELECT id FROM chunks WHERE file_id=?", (file_id,)).fetchall()]
        quest_ids = [item["id"] for item in self.conn.execute("SELECT id FROM quests WHERE file_id=?", (file_id,)).fetchall()]
        self.conn.execute("DELETE FROM files WHERE id=?", (file_id,))
        self.conn.execute("DELETE FROM chunks WHERE file_id=?", (file_id,))
        self.conn.execute("DELETE FROM quests WHERE file_id=?", (file_id,))
        if self.fts_enabled:
            self.conn.execute("DELETE FROM chunk_fts WHERE file_id=?", (file_id,))
        self.conn.execute("DELETE FROM embeddings WHERE owner_type='doc' AND owner_id=?", (file_id,))
        for chunk_id in chunk_ids:
            self.conn.execute("DELETE FROM embeddings WHERE owner_type='chunk' AND owner_id=?", (chunk_id,))
        for quest_id in quest_ids:
            self.conn.execute("DELETE FROM embeddings WHERE owner_type='quest' AND owner_id=?", (quest_id,))
        self.conn.commit()
        try:
            if mirror_path.exists() and TREE_DIR.resolve() in mirror_path.resolve().parents:
                mirror_path.unlink()
        except Exception:
            pass
        try:
            artifact_dir = artifact_path.parent if artifact_path.name else Path("")
            if artifact_dir.exists() and ARTIFACTS_DIR.resolve() in artifact_dir.resolve().parents:
                shutil.rmtree(artifact_dir, ignore_errors=True)
        except Exception:
            pass
        self._invalidate_cache()
        self._update_collection_counts(collection_id)
        return True

    def delete_collection(self, collection_id: str) -> bool:
        collection = self.conn.execute("SELECT id FROM collections WHERE id=?", (collection_id,)).fetchone()
        rows = self.conn.execute("SELECT id FROM files WHERE collection_id=?", (collection_id,)).fetchall()
        if not collection and not rows:
            return False
        for row in rows:
            self.delete_document(row["id"])
        self.conn.execute("DELETE FROM jobs WHERE collection_id=?", (collection_id,))
        self.conn.execute("DELETE FROM collections WHERE id=?", (collection_id,))
        self.conn.commit()
        self._invalidate_cache()
        return True

    def delete_folder(self, folder_path: str) -> bool:
        folder = _safe_relative_path(folder_path or "")
        if not folder:
            return False
        rows = self.conn.execute(
            "SELECT id FROM files WHERE relative_path=? OR relative_path LIKE ?",
            (folder, folder.rstrip("/") + "/%"),
        ).fetchall()
        if not rows:
            return False
        for row in rows:
            self.delete_document(row["id"])
        try:
            tree_dir = TREE_DIR / folder
            if tree_dir.exists() and TREE_DIR.resolve() in tree_dir.resolve().parents:
                shutil.rmtree(tree_dir, ignore_errors=True)
        except Exception:
            pass
        return True

    def expand_archive(self, filename: str, content: bytes, relative_path: Optional[str] = None) -> list[dict]:
        if Path(filename).suffix.lower() not in ARCHIVE_EXTS:
            return [{"filename": filename, "relative_path": relative_path or filename, "content": content}]
        files = []
        base = PurePosixPath(relative_path or filename).stem
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for info in archive.infolist():
                if info.is_dir() or info.file_size == 0:
                    continue
                inner = _safe_relative_path(info.filename)
                files.append({"filename": PurePosixPath(inner).name, "relative_path": f"{base}/{inner}", "content": archive.read(info)})
        return files

    def _normalize_scopes(self, scopes: Optional[list[dict]]) -> Optional[dict[str, set[str]]]:
        if not scopes:
            return None
        normalized = {"collection": set(), "folder": set(), "file": set()}
        for scope in scopes:
            if not isinstance(scope, dict):
                continue
            kind = scope.get("type") or scope.get("kind")
            value = str(scope.get("id") or scope.get("path") or scope.get("value") or "").strip()
            if value and kind in normalized:
                normalized[kind].add(_safe_relative_path(value) if kind == "folder" else value)
        return normalized if any(normalized.values()) else None

    def _file_in_scope(self, file_row, scope: Optional[dict[str, set[str]]]) -> bool:
        if not scope:
            return True
        if file_row["id"] in scope["file"] or file_row["collection_id"] in scope["collection"]:
            return True
        path = (file_row["relative_path"] or file_row["filename"] or "").strip("/")
        return any(path == folder.strip("/") or path.startswith(folder.strip("/") + "/") for folder in scope["folder"])

    def _fts_search(self, query: str, top_k: int, scope: Optional[dict[str, set[str]]]):
        if not self.fts_enabled:
            return []
        fts_query, coverage = self._fts_query_parts(query)
        if not fts_query:
            return []
        try:
            rows = self.conn.execute(
                """SELECT chunk_id, file_id, bm25(chunk_fts) AS rank
                   FROM chunk_fts
                   WHERE chunk_fts MATCH ?
                   ORDER BY rank
                   LIMIT ?""",
                (fts_query, max(top_k * 4, 25)),
            ).fetchall()
        except sqlite3.Error as exc:
            print(f"[DOCS] FTS query failed: {exc}")
            return []
        hits = []
        for index, row in enumerate(rows):
            file_row = self._get_file(row["file_id"])
            if not file_row or not self._file_in_scope(file_row, scope):
                continue
            rank_strength = min(1.0, math.log1p(max(0.0, -float(row["rank"]))) / math.log1p(12.0))
            score = max(0.12, 0.2 + (0.65 * rank_strength * coverage) - (index * 0.01))
            hits.append({"id": row["chunk_id"], "score": score})
            if len(hits) >= top_k:
                break
        return hits

    def _fts_query(self, query: str) -> str:
        return self._fts_query_parts(query)[0]

    def _fts_query_parts(self, query: str) -> tuple[str, float]:
        seen = set()
        tokens = []
        for token in re.findall(r"[A-Za-z0-9\u0900-\u0d7f]+", (query or "").lower()):
            if len(token) <= 2 or token in seen:
                continue
            seen.add(token)
            tokens.append(token)
            if len(tokens) >= 24:
                break
        if not tokens:
            return "", 0.0
        total = self.conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] or 1
        broad_cutoff = max(20, int(total * 0.18))
        stats = []
        absent = 0
        for token in tokens:
            try:
                count = self.conn.execute(
                    "SELECT COUNT(*) FROM chunk_fts WHERE chunk_fts MATCH ?",
                    (f'"{token}"',),
                ).fetchone()[0]
            except sqlite3.Error:
                continue
            if count:
                stats.append((token, count))
            elif len(token) > 3:
                absent += 1
        if not stats:
            return "", 0.0
        selected = [token for token, count in stats if count <= broad_cutoff]
        if not selected:
            selected = [token for token, _ in sorted(stats, key=lambda item: item[1])[:4]]
        selected = selected[:12]
        coverage = len(selected) / max(1, len(selected) + absent)
        return " OR ".join(f'"{token}"' for token in selected), coverage

    def _search(self, owner_type: str, q_vec, top_k: int, scope: Optional[dict[str, set[str]]]):
        ids, matrix = self._load_vectors(owner_type)
        if not ids or matrix.size == 0:
            return []
        sims = matrix @ q_vec.astype(np.float32)
        hits = []
        for index in np.argsort(sims)[::-1]:
            score = float(sims[index])
            if score <= 0.05:
                break
            owner_id = ids[int(index)]
            file_row = self._get_file(self._owner_file_id(owner_type, owner_id))
            if file_row and self._file_in_scope(file_row, scope):
                hits.append({"id": owner_id, "score": score})
            if len(hits) >= top_k:
                break
        return hits

    def _owner_file_id(self, owner_type: str, owner_id: str) -> Optional[str]:
        if owner_type == "doc":
            return owner_id
        table = "chunks" if owner_type == "chunk" else "quests"
        row = self.conn.execute(f"SELECT file_id FROM {table} WHERE id=?", (owner_id,)).fetchone()
        return row["file_id"] if row else None

    def _put_embedding(self, owner_type: str, owner_id: str, vector):
        vector = np.asarray(vector, dtype=np.float32)
        self.conn.execute(
            "INSERT OR REPLACE INTO embeddings(owner_type,owner_id,dim,vector) VALUES(?,?,?,?)",
            (owner_type, owner_id, int(vector.shape[0]), vector.tobytes()),
        )

    def _load_vectors(self, owner_type: str):
        if owner_type in self._cache:
            return self._cache[owner_type]
        rows = self.conn.execute("SELECT owner_id,dim,vector FROM embeddings WHERE owner_type=?", (owner_type,)).fetchall()
        ids, vectors = [], []
        for row in rows:
            vector = np.frombuffer(row["vector"], dtype=np.float32)
            if vector.shape[0] == row["dim"]:
                ids.append(row["owner_id"])
                vectors.append(vector)
        matrix = np.vstack(vectors) if vectors else np.zeros((0, EMBED_DIM), dtype=np.float32)
        self._cache[owner_type] = (ids, matrix)
        return ids, matrix

    def _invalidate_cache(self):
        self._cache = {}

    def _merge_candidate(self, candidates, key, score, text, file_id, result_type, prefix, chunk):
        file_row = self._get_file(file_id)
        if not file_row or not text:
            return
        existing = candidates.get(key)
        if existing and existing["score"] >= score:
            return
        source_path = chunk["source_path"] if chunk else file_row["relative_path"]
        source_bits = [source_path]
        if chunk and chunk["page"]:
            source_bits.append(f"page {chunk['page']}")
        if chunk and chunk["sheet"]:
            source_bits.append(f"sheet {chunk['sheet']}")
        source = " | ".join(str(value) for value in source_bits if value)
        candidates[key] = {
            "text": f"Source: {source}\n{prefix}{text}".strip(),
            "score": float(score),
            "doc_id": file_id,
            "file_id": file_id,
            "type": result_type,
            "source_path": source_path,
            "page": chunk["page"] if chunk else None,
            "sheet": chunk["sheet"] if chunk else None,
        }

    def _get_file(self, file_id: Optional[str]):
        if not file_id:
            return None
        return self.conn.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()

    def _get_chunk(self, chunk_id: str):
        return self.conn.execute("SELECT * FROM chunks WHERE id=?", (chunk_id,)).fetchone()

    def _get_quest(self, quest_id: str):
        return self.conn.execute("SELECT * FROM quests WHERE id=?", (quest_id,)).fetchone()

    def _first_chunk(self, file_id: str):
        return self.conn.execute("SELECT * FROM chunks WHERE file_id=? ORDER BY chunk_index LIMIT 1", (file_id,)).fetchone()

    def _file_to_dict(self, row) -> dict:
        keys = row.keys()
        return {
            "id": row["id"],
            "filename": row["filename"],
            "relative_path": row["relative_path"],
            "collection_id": row["collection_id"],
            "collection_name": row["collection_name"] if "collection_name" in keys else None,
            "mime_type": row["mime_type"],
            "extension": row["extension"],
            "size_bytes": row["size_bytes"],
            "parser": row["parser"],
            "status": row["status"],
            "error": row["error"],
            "chunk_count": row["chunk_count"],
            "quest_count": row["quest_count"],
            "quest_status": row["quest_status"],
            "faq_status": row["quest_status"],
            "char_count": row["char_count"],
            "token_count": row["token_count"],
        }

    def _tree_to_public(self, node: dict) -> dict:
        children = [self._tree_to_public(child) for child in sorted(node["children"].values(), key=lambda item: item["name"].lower())]
        files = sorted(node["files"], key=lambda item: (item.get("relative_path") or item.get("filename") or "").lower())
        return {"name": node["name"], "path": node["path"], "type": node["type"], "children": children, "files": files}

    def _update_collection_counts(self, collection_id: str):
        rows = self.conn.execute("SELECT status, COUNT(*) AS n FROM files WHERE collection_id=? GROUP BY status", (collection_id,)).fetchall()
        counts = {row["status"]: row["n"] for row in rows}
        total = sum(counts.values())
        ready = counts.get("ready", 0)
        failed = counts.get("failed", 0)
        status = "ready" if total and ready + failed == total else "ingesting"
        if total == 0:
            status = "empty"
        self.conn.execute(
            "UPDATE collections SET total_files=?,ready_files=?,failed_files=?,status=?,updated_at=? WHERE id=?",
            (total, ready, failed, status, _now(), collection_id),
        )
        self.conn.commit()

    def _document_embedding_text(self, filename: str, rel_path: str, full_text: str, chunks: list[dict]) -> str:
        headings = [chunk["text"].splitlines()[0] for chunk in chunks[:8] if chunk.get("text")]
        return f"Document: {rel_path or filename}\n" + "\n".join(headings) + "\n\n" + full_text[:6000]

    def _quest_target_count(self, token_count: int) -> int:
        return max(5, min(150, math.ceil(max(token_count, 1) / 500)))

    def _select_quest_chunks(self, chunks, target_k: int):
        limit = max(2, min(len(chunks), math.ceil(target_k / 3)))
        scored = []
        for row in chunks:
            text = row["text"] or ""
            score = min(len(text), 2500) / 2500
            scored.append((score, row))
        chosen = [row for _, row in sorted(scored, key=lambda item: item[0], reverse=True)[:limit]]
        chosen.sort(key=lambda row: row["chunk_index"])
        return chosen

    def _parse_questions(self, text: str) -> list[str]:
        questions = []
        for line in (text or "").splitlines():
            line = re.sub(r"^[-*\d.)\s]+", "", line).strip().strip('"')
            if line:
                questions.append(line[:line.rfind("?") + 1] if "?" in line else line)
        return questions

    def _clean_question(self, question: str) -> str:
        question = re.sub(r"\s+", " ", question or "").strip(" -\t\n\r")
        if len(question) < 12 or len(question) > 240:
            return ""
        return question if question.endswith("?") else question + "?"

    def _fallback_questions(self, file_row, chunk, count: int) -> list[str]:
        title = Path(file_row["relative_path"] or file_row["filename"]).stem.replace("_", " ").replace("-", " ")
        phrases = []
        for line in (chunk["text"] or "").splitlines():
            line = re.sub(r"[#>*`|]+", "", line).strip()
            if 4 <= len(line.split()) <= 14:
                phrases.append(line)
            if len(phrases) >= count:
                break
        while len(phrases) < count:
            phrases.append(title)
        return [f"What does {title} say about {phrase}?" for phrase in phrases[:count]]

    def _parse_file(self, filename: str, rel_path: str, content: bytes) -> dict:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf(rel_path, content)
        if ext == ".docx":
            return self._parse_docx(rel_path, content)
        if ext in {".csv", ".tsv"}:
            return self._parse_csv(rel_path, content, "\t" if ext == ".tsv" else ",")
        if ext in {".xlsx", ".xlsm"}:
            return self._parse_xlsx(rel_path, content)
        if ext in {".html", ".htm"}:
            return self._parse_html(rel_path, content)
        if ext in IMAGE_EXTS:
            return self._parse_image(rel_path, content)
        return self._parse_text(rel_path, content, parser="text")

    def _parse_pdf(self, rel_path: str, content: bytes) -> dict:
        from pypdf import PdfReader
        elements = []
        for page_index, page in enumerate(PdfReader(io.BytesIO(content)).pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                elements.append({"type": "page", "text": text.strip(), "source_path": rel_path, "page": page_index})
        return {"parser": "pypdf", "elements": elements}

    def _parse_docx(self, rel_path: str, content: bytes) -> dict:
        import docx
        doc = docx.Document(io.BytesIO(content))
        elements = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                style = (para.style.name if para.style else "").lower()
                elements.append({"type": "heading" if "heading" in style else "paragraph", "text": text, "source_path": rel_path})
        for table_index, table in enumerate(doc.tables, start=1):
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            text = self._rows_to_markdown(rows)
            if text.strip():
                elements.append({"type": "table", "text": text, "source_path": rel_path, "metadata": {"table": table_index}})
        return {"parser": "python-docx", "elements": elements}

    def _parse_csv(self, rel_path: str, content: bytes, delimiter: str) -> dict:
        text = _decode_text(content)
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except Exception:
            dialect = csv.excel_tab if delimiter == "\t" else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        if not rows:
            return {"parser": "csv", "elements": []}
        elements, header, start, batch = [], rows[0], 2, []
        for row_index, row in enumerate(rows[1:], start=2):
            batch.append(row)
            if len(batch) >= 40:
                elements.append({"type": "table", "text": self._rows_to_markdown([header] + batch), "source_path": rel_path, "row_start": start, "row_end": row_index})
                start, batch = row_index + 1, []
        if batch:
            elements.append({"type": "table", "text": self._rows_to_markdown([header] + batch), "source_path": rel_path, "row_start": start, "row_end": len(rows)})
        if not elements:
            elements.append({"type": "table", "text": self._rows_to_markdown(rows), "source_path": rel_path, "row_start": 1, "row_end": len(rows)})
        return {"parser": "csv", "elements": elements}

    def _parse_xlsx(self, rel_path: str, content: bytes) -> dict:
        elements = []
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            shared = self._xlsx_shared_strings(archive)
            sheet_names = self._xlsx_sheet_names(archive)
            sheet_files = sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
            for sheet_index, sheet_file in enumerate(sheet_files, start=1):
                rows = self._xlsx_rows(archive, sheet_file, shared)
                if not rows:
                    continue
                sheet_name = sheet_names.get(sheet_index, f"Sheet {sheet_index}")
                header, start, batch = rows[0], 2, []
                for row_index, row in enumerate(rows[1:], start=2):
                    batch.append(row)
                    if len(batch) >= 35:
                        elements.append({"type": "sheet", "text": self._rows_to_markdown([header] + batch), "source_path": rel_path, "sheet": sheet_name, "row_start": start, "row_end": row_index})
                        start, batch = row_index + 1, []
                if batch:
                    elements.append({"type": "sheet", "text": self._rows_to_markdown([header] + batch), "source_path": rel_path, "sheet": sheet_name, "row_start": start, "row_end": len(rows)})
                elif rows:
                    elements.append({"type": "sheet", "text": self._rows_to_markdown(rows), "source_path": rel_path, "sheet": sheet_name, "row_start": 1, "row_end": len(rows)})
        return {"parser": "xlsx-xml", "elements": elements}

    def _parse_html(self, rel_path: str, content: bytes) -> dict:
        text = _decode_text(content)
        try:
            from lxml import html
            root = html.fromstring(text)
            for bad in root.xpath("//script|//style|//noscript"):
                bad.drop_tree()
            extracted = root.text_content()
        except Exception:
            extracted = re.sub(r"<[^>]+>", " ", text)
        return {"parser": "html", "elements": [{"type": "html", "text": re.sub(r"\n{3,}", "\n\n", extracted).strip(), "source_path": rel_path}]}

    def _parse_image(self, rel_path: str, content: bytes) -> dict:
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(content))
            meta = f"Image file {rel_path}. Format: {image.format}. Size: {image.width}x{image.height}."
            try:
                import pytesseract
                ocr_text = pytesseract.image_to_string(image).strip()
                parser = "pillow+pytesseract"
                text = meta + ("\n\nOCR text:\n" + ocr_text if ocr_text else "")
            except Exception:
                parser = "pillow-metadata"
                text = meta + " OCR is not installed, so visual text could not be extracted."
            return {"parser": parser, "elements": [{"type": "image", "text": text, "source_path": rel_path}]}
        except Exception as exc:
            return {"parser": "image", "elements": [], "error": str(exc)}

    def _parse_text(self, rel_path: str, content: bytes, parser: str) -> dict:
        text = _decode_text(content)
        if Path(rel_path).suffix.lower() == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except Exception:
                pass
        return {"parser": parser, "elements": [{"type": "text", "text": text.strip(), "source_path": rel_path}]}

    def _chunk_elements(self, elements: list[dict], max_chars: int = 1800, overlap_chars: int = 250) -> list[dict]:
        chunks, current, current_meta, current_len = [], [], {}, 0

        def flush():
            nonlocal current, current_meta, current_len
            text = "\n\n".join(current).strip()
            if text:
                chunk = dict(current_meta)
                chunk["text"] = text
                chunk["type"] = chunk.get("type", "chunk")
                chunks.append(chunk)
            current = [text[-overlap_chars:]] if overlap_chars and text else []
            current_len = len(current[0]) if current else 0

        for element in elements:
            text = (element.get("text") or "").strip()
            if not text:
                continue
            label = element.get("type", "text")
            if label == "heading":
                text = f"# {text}"
            elif label in {"table", "sheet"}:
                text = f"Table from {element.get('sheet') or element.get('source_path')}:\n{text}"
            meta = {
                "source_path": element.get("source_path"),
                "page": element.get("page"),
                "sheet": element.get("sheet"),
                "row_start": element.get("row_start"),
                "row_end": element.get("row_end"),
                "metadata": element.get("metadata", {}),
                "type": label if label in {"table", "sheet", "image"} else "chunk",
            }
            if current and current_len + len(text) + 2 > max_chars:
                flush()
            if not current:
                current_meta = meta
            if len(text) > max_chars * 1.5:
                step = max_chars - overlap_chars
                for index in range(0, len(text), step):
                    part = text[index:index + max_chars].strip()
                    if part:
                        chunk = dict(meta)
                        chunk["text"] = part
                        chunks.append(chunk)
                current, current_len = [], 0
                continue
            current.append(text)
            current_len += len(text) + 2
        if current:
            flush()
        return chunks

    def _rows_to_markdown(self, rows: list[list[str]]) -> str:
        rows = [[str(cell or "").strip() for cell in row] for row in rows if any(str(cell or "").strip() for cell in row)]
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        lines = ["| " + " | ".join(cell.replace("|", "\\|") for cell in rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
        for row in rows[1:]:
            lines.append("| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |")
        return "\n".join(lines)

    def _xlsx_shared_strings(self, archive) -> list[str]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        return ["".join(text.text or "" for text in item.iter(f"{ns}t")) for item in root.findall(f"{ns}si")]

    def _xlsx_sheet_names(self, archive) -> dict[int, str]:
        if "xl/workbook.xml" not in archive.namelist():
            return {}
        root = ET.fromstring(archive.read("xl/workbook.xml"))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        return {index: sheet.attrib.get("name", f"Sheet {index}") for index, sheet in enumerate(root.findall(f".//{ns}sheet"), start=1)}

    def _xlsx_rows(self, archive, sheet_file: str, shared: list[str]) -> list[list[str]]:
        root = ET.fromstring(archive.read(sheet_file))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        rows = []
        for row in root.findall(f".//{ns}row"):
            values = []
            for cell in row.findall(f"{ns}c"):
                ctype = cell.attrib.get("t")
                value_element = cell.find(f"{ns}v")
                value = ""
                if ctype == "inlineStr":
                    value = "".join(text.text or "" for text in cell.iter(f"{ns}t"))
                elif value_element is not None and value_element.text is not None:
                    raw = value_element.text
                    if ctype == "s":
                        try:
                            value = shared[int(raw)]
                        except Exception:
                            value = raw
                    else:
                        value = raw
                values.append(value)
            if any(value.strip() for value in values):
                rows.append(values)
        return rows


class RAGPipeline(DocumentStore):
    pass


# Shared document store used by the voice loop and the REST API.
rag = RAGPipeline()
